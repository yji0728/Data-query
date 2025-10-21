"""
Excel Data Extractor
Extracts data from Excel files using openpyxl and pandas.
"""

import pandas as pd
import openpyxl
from typing import List, Dict, Any, Optional, Union
from pathlib import Path
import logging


class ExcelExtractor:
    """Extracts data from Excel files."""
    
    def __init__(self):
        """Initialize Excel extractor."""
        self.logger = logging.getLogger(__name__)
        
    def read_excel_pandas(self, file_path: str, sheet_name: Optional[Union[str, int]] = 0,
                         header: Optional[int] = 0, 
                         usecols: Optional[List[str]] = None) -> List[Dict[str, Any]]:
        """
        Read Excel file using pandas.
        
        Args:
            file_path: Path to Excel file
            sheet_name: Sheet name or index (default: first sheet)
            header: Row to use as column names (default: 0)
            usecols: List of column names to read
            
        Returns:
            List of dictionaries containing row data
        """
        try:
            df = pd.read_excel(file_path, sheet_name=sheet_name, 
                             header=header, usecols=usecols)
            
            # Convert to list of dictionaries
            data = df.to_dict('records')
            
            # Clean NaN values
            for row in data:
                for key, value in row.items():
                    if pd.isna(value):
                        row[key] = None
                        
            self.logger.info(f"Read {len(data)} rows from {file_path}")
            return data
        except Exception as e:
            self.logger.error(f"Error reading Excel file with pandas: {e}")
            raise
            
    def read_excel_openpyxl(self, file_path: str, 
                           sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
        """
        Read Excel file using openpyxl.
        
        Args:
            file_path: Path to Excel file
            sheet_name: Sheet name (default: active sheet)
            
        Returns:
            List of dictionaries containing row data
        """
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            
            if sheet_name:
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active
                
            # Get headers from first row
            headers = []
            for cell in sheet[1]:
                headers.append(cell.value)
                
            # Extract data rows
            data = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(row):  # Skip empty rows
                    row_data = {headers[i]: row[i] for i in range(len(headers))}
                    data.append(row_data)
                    
            workbook.close()
            self.logger.info(f"Read {len(data)} rows from {file_path}")
            return data
        except Exception as e:
            self.logger.error(f"Error reading Excel file with openpyxl: {e}")
            raise
            
    def get_sheet_names(self, file_path: str) -> List[str]:
        """
        Get list of sheet names in an Excel file.
        
        Args:
            file_path: Path to Excel file
            
        Returns:
            List of sheet names
        """
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            return sheet_names
        except Exception as e:
            self.logger.error(f"Error getting sheet names: {e}")
            raise
            
    def read_specific_range(self, file_path: str, sheet_name: Optional[str] = None,
                           min_row: int = 1, max_row: Optional[int] = None,
                           min_col: int = 1, max_col: Optional[int] = None) -> List[List[Any]]:
        """
        Read a specific range from an Excel sheet.
        
        Args:
            file_path: Path to Excel file
            sheet_name: Sheet name (default: active sheet)
            min_row: Minimum row number
            max_row: Maximum row number (optional)
            min_col: Minimum column number
            max_col: Maximum column number (optional)
            
        Returns:
            List of lists containing cell values
        """
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            
            if sheet_name:
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active
                
            data = []
            for row in sheet.iter_rows(min_row=min_row, max_row=max_row,
                                      min_col=min_col, max_col=max_col,
                                      values_only=True):
                data.append(list(row))
                
            workbook.close()
            self.logger.info(f"Read {len(data)} rows from range")
            return data
        except Exception as e:
            self.logger.error(f"Error reading specific range: {e}")
            raise
            
    def read_all_sheets(self, file_path: str) -> Dict[str, List[Dict[str, Any]]]:
        """
        Read all sheets from an Excel file.
        
        Args:
            file_path: Path to Excel file
            
        Returns:
            Dictionary with sheet names as keys and data as values
        """
        try:
            all_data = {}
            sheet_names = self.get_sheet_names(file_path)
            
            for sheet_name in sheet_names:
                data = self.read_excel_pandas(file_path, sheet_name=sheet_name)
                all_data[sheet_name] = data
                
            self.logger.info(f"Read {len(all_data)} sheets from {file_path}")
            return all_data
        except Exception as e:
            self.logger.error(f"Error reading all sheets: {e}")
            raise
            
    def filter_data(self, data: List[Dict[str, Any]], 
                   filters: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        Filter data based on conditions.
        
        Args:
            data: List of dictionaries to filter
            filters: Dictionary of column names and filter values
            
        Returns:
            Filtered list of dictionaries
        """
        filtered_data = []
        
        for row in data:
            match = True
            for col, value in filters.items():
                if col not in row or row[col] != value:
                    match = False
                    break
            if match:
                filtered_data.append(row)
                
        self.logger.info(f"Filtered to {len(filtered_data)} rows")
        return filtered_data
        
    def select_columns(self, data: List[Dict[str, Any]], 
                      columns: List[str]) -> List[Dict[str, Any]]:
        """
        Select specific columns from data.
        
        Args:
            data: List of dictionaries
            columns: List of column names to select
            
        Returns:
            List of dictionaries with only selected columns
        """
        selected_data = []
        
        for row in data:
            selected_row = {col: row.get(col) for col in columns}
            selected_data.append(selected_row)
            
        return selected_data
        
    def transform_data(self, data: List[Dict[str, Any]], 
                      transformations: Dict[str, callable]) -> List[Dict[str, Any]]:
        """
        Apply transformations to data.
        
        Args:
            data: List of dictionaries to transform
            transformations: Dictionary mapping column names to transformation functions
            
        Returns:
            Transformed data
        """
        transformed_data = []
        
        for row in data:
            new_row = row.copy()
            for col, transform_func in transformations.items():
                if col in new_row:
                    try:
                        new_row[col] = transform_func(new_row[col])
                    except Exception as e:
                        self.logger.warning(f"Error transforming column {col}: {e}")
            transformed_data.append(new_row)
            
        return transformed_data
